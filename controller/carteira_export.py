"""
Exportação da carteira do mês (Relatórios → Exportar Carteira).

Permissões:
- consultor: só a própria carteira
- gestão: a carteira dos executivos da sua equipe no mês
- planejamento e admin: todos

Premissas (as mesmas do Gerenciador de Carteira):
- tbcarteira(periodo, vendedor, id_condominio), periodo = 1º dia do mês
- tbhierarquia(id_usuario, gestor_direto, periodo) define a equipe
"""
import re
from datetime import date
from io import BytesIO

import pandas as pd

from controller.carteira_control import _consulta

COLUNAS = ["id_condominio", "nome_condominio", "endereco", "executivo", "gestor_direto"]
TITULOS = {
    "id_condominio": "ID Condomínio",
    "nome_condominio": "Condomínio",
    "endereco": "Endereço completo",
    "executivo": "Executivo",
    "gestor_direto": "Gestor direto",
}
MESES = ["Jan", "Fev", "Mar", "Abr", "Mai", "Jun", "Jul", "Ago", "Set", "Out", "Nov", "Dez"]

# Azul-noite e verde da Alloha
AZUL, VERDE, CINZA = "001059", "44EE67", "F2F4FA"


def mes_atual_label():
    hoje = date.today()
    return f"{MESES[hoje.month - 1]}/{str(hoje.year)[2:]}"


def _cep(v):
    d = re.sub(r"\D", "", str(v or ""))
    return f"{d[:5]}-{d[5:]}" if len(d) == 8 else (v or "")


def _carteira_bruta(role, username, executivo=None, gestor=None):
    """Carteira do mês atual conforme o perfil, com as chaves (username) de executivo e gestor.
    executivo/gestor são USERNAMES, não nomes de exibição."""
    where, params = [], []
    if role == "consultor":
        where.append("lower(k.vendedor) = %s")
        params.append(username.lower())
    elif role == "gestao":
        where.append("lower(h.gestor_direto) = %s")
        params.append(username.lower())
    if executivo:
        where.append("lower(k.vendedor) = %s")
        params.append(executivo.lower())
    if gestor:
        where.append("lower(h.gestor_direto) = %s")
        params.append(gestor.lower())

    filtro = (" AND " + " AND ".join(where)) if where else ""
    df = _consulta(
        f"""
        SELECT k.id_condominio,
               c.nome  AS nome_condominio,
               c.cidade, c.sigla_estado, c.bairro, c.logradouro, c.numero, c.cep,
               COALESCE(u.nome, k.vendedor) AS executivo,
               COALESCE(g.nome, h.gestor_direto, '') AS gestor_direto,
               lower(k.vendedor) AS executivo_key,
               lower(COALESCE(h.gestor_direto, '')) AS gestor_key
        FROM tbcarteira k
        LEFT JOIN tb_condominio c ON c.id::text = k.id_condominio::text
        LEFT JOIN tbusuarios u ON lower(u.username) = lower(k.vendedor)
        LEFT JOIN tbhierarquia h ON h.id_usuario = u.id
             AND date_trunc('month', h.periodo::date) = date_trunc('month', current_date)
        LEFT JOIN tbusuarios g ON lower(g.username) = lower(h.gestor_direto)
        WHERE date_trunc('month', k.periodo::date) = date_trunc('month', current_date){filtro}
        ORDER BY executivo, c.nome;
        """,
        tuple(params),
    )
    return df


def carteira_para_export(role, username, executivo=None, gestor=None):
    """Retorna a carteira do mês atual conforme o perfil, já formatada."""
    df = _carteira_bruta(role, username, executivo, gestor)
    if df.empty:
        return pd.DataFrame(columns=COLUNAS)

    partes = df[["cidade", "sigla_estado", "bairro", "logradouro", "numero", "cep"]].fillna("").astype(str)
    df["endereco"] = [
        ", ".join(
            p for p in [
                f"{r.cidade}-{r.sigla_estado}".strip("-"), r.bairro, r.logradouro, r.numero, _cep(r.cep)
            ] if p.strip()
        )
        for r in partes.itertuples()
    ]
    return df[COLUNAS].fillna("")


def opcoes_filtro(role, username):
    """Executivos e gestores da carteira do mês, como {username: nome de exibição}.
    O filtro usa o username (chave do banco); o nome é só para mostrar no selectbox."""
    df = _carteira_bruta(role, username)
    if df.empty:
        return {}, {}
    execs = dict(sorted(
        df[["executivo_key", "executivo"]].drop_duplicates("executivo_key").itertuples(index=False),
        key=lambda kv: str(kv[1]).lower(),
    ))
    g = df[df.gestor_key != ""][["gestor_key", "gestor_direto"]].drop_duplicates("gestor_key")
    gestores = dict(sorted(g.itertuples(index=False), key=lambda kv: str(kv[1]).lower()))
    return execs, gestores


# ── Arquivos ──────────────────────────────────────────────────────────────
def para_excel(df, titulo):
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    saida = BytesIO()
    dados = df.rename(columns=TITULOS)
    with pd.ExcelWriter(saida, engine="openpyxl") as w:
        dados.to_excel(w, index=False, sheet_name="Carteira", startrow=2)
        ws = w.sheets["Carteira"]
        ws.cell(row=1, column=1, value=titulo).font = Font(bold=True, size=14, color=AZUL)
        cabecalho = PatternFill("solid", fgColor=AZUL)
        for col in range(1, len(dados.columns) + 1):
            c = ws.cell(row=3, column=col)
            c.fill, c.font = cabecalho, Font(bold=True, color="FFFFFF")
            c.alignment = Alignment(vertical="center")
        larguras = {"ID Condomínio": 26, "Condomínio": 38, "Endereço completo": 62,
                    "Executivo": 24, "Gestor direto": 24}
        for i, nome in enumerate(dados.columns, start=1):
            ws.column_dimensions[get_column_letter(i)].width = larguras.get(nome, 20)
        ws.freeze_panes = "A4"
        ws.auto_filter.ref = f"A3:{get_column_letter(len(dados.columns))}{len(dados) + 3}"
    return saida.getvalue()


def para_pdf(df, titulo):
    from reportlab.lib import colors
    from reportlab.lib.enums import TA_LEFT
    from reportlab.lib.pagesizes import A4, landscape
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.lib.units import mm
    from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

    azul, verde = colors.HexColor("#" + AZUL), colors.HexColor("#" + VERDE)
    saida = BytesIO()
    doc = SimpleDocTemplate(
        saida, pagesize=landscape(A4), title=titulo,
        leftMargin=12 * mm, rightMargin=12 * mm, topMargin=12 * mm, bottomMargin=14 * mm,
    )
    h1 = ParagraphStyle("h1", fontName="Helvetica-Bold", fontSize=15, textColor=azul, alignment=TA_LEFT)
    sub = ParagraphStyle("sub", fontName="Helvetica", fontSize=9, textColor=colors.HexColor("#5A6391"))
    cel = ParagraphStyle("cel", fontName="Helvetica", fontSize=7.5, leading=9.5)
    celb = ParagraphStyle("celb", parent=cel, fontName="Helvetica-Bold", textColor=colors.white)

    linhas = [[Paragraph(TITULOS[c], celb) for c in COLUNAS]]
    linhas += [[Paragraph(str(v), cel) for v in r] for r in df[COLUNAS].values.tolist()]

    tab = Table(linhas, repeatRows=1, colWidths=[42 * mm, 62 * mm, 98 * mm, 33 * mm, 33 * mm])
    tab.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), azul),
        ("LINEBELOW", (0, 0), (-1, 0), 1.2, verde),
        ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#" + CINZA)]),
        ("GRID", (0, 0), (-1, -1), 0.25, colors.HexColor("#D6DCEB")),
        ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
        ("LEFTPADDING", (0, 0), (-1, -1), 4),
        ("RIGHTPADDING", (0, 0), (-1, -1), 4),
        ("TOPPADDING", (0, 0), (-1, -1), 3),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 3),
    ]))

    def rodape(canvas, documento):
        canvas.saveState()
        canvas.setFont("Helvetica", 7.5)
        canvas.setFillColor(colors.HexColor("#5A6391"))
        canvas.drawString(12 * mm, 8 * mm, "Portal Alto Valor · Giga+ Fibra / Alloha Fibra")
        canvas.drawRightString(landscape(A4)[0] - 12 * mm, 8 * mm, f"Página {documento.page}")
        canvas.restoreState()

    doc.build(
        [Paragraph(titulo, h1), Spacer(1, 3), Paragraph(f"{len(df)} prédio(s) na carteira", sub), Spacer(1, 8), tab],
        onFirstPage=rodape, onLaterPages=rodape,
    )
    return saida.getvalue()
